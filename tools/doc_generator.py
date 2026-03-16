"""Document generator tool - creates .docx requirements document from complete_dataset.json."""
import json
import os
import sys
from datetime import datetime
from pathlib import Path

from crewai.tools import BaseTool

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import get_config
from tools.data_output import get_output_dir

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _add_paragraph_shading(paragraph, fill_color):
    """Add background color to a paragraph."""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), fill_color)
    paragraph._element.get_or_add_pPr().append(shading_elm)


def _create_bullet_paragraph(doc, text, level=0, bold_prefix=None):
    """Create a bullet paragraph at the given indentation level."""
    bullet_chars = ['•', '○', '▪']
    indents = [0.5, 1.0, 1.5]

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indents[level])
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    run = p.add_run(f"{bullet_chars[level]} ")
    run.font.name = 'Arial'
    run.font.size = Pt(11)

    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.bold = True

    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)

    return p


def generate_requirements_doc(data: dict, output_path: str, model_name: str = "CrewAI Agents") -> str:
    """Generate a .docx requirements document from complete dataset."""
    doc = Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    title = doc.add_heading('Test Scenarios - MIRA Multi-Agent System', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_paragraph_shading(title, '1F1F1F')
    for run in title.runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.name = 'Arial'
        run.font.size = Pt(28)
        run.font.bold = True

    # Metadata
    now = datetime.now()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)

    run = p.add_run('Model: ')
    run.font.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(11)

    run = p.add_run(f'{model_name} ')
    run.font.name = 'Arial'
    run.font.size = Pt(11)

    run = p.add_run('Generated on: ')
    run.font.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(11)

    run = p.add_run(now.strftime('%Y-%m-%d %H:%M:%S'))
    run.font.name = 'Arial'
    run.font.size = Pt(11)

    # Separator
    p = doc.add_paragraph('_' * 79)
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(11)

    # Intro
    p = doc.add_paragraph(
        'Here are the test scenarios for each non-functional requirement, '
        'covering various conditions and checkpoints:'
    )
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(20)
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(11)

    p = doc.add_paragraph('_' * 79)
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(11)

    # Group data — handle items that may be strings instead of dicts
    def _ensure_dict_list(items):
        """Ensure all items in a list are dicts."""
        result = []
        for item in items:
            if isinstance(item, dict):
                result.append(item)
            elif isinstance(item, str):
                try:
                    parsed = json.loads(item)
                    if isinstance(parsed, dict):
                        result.append(parsed)
                    elif isinstance(parsed, list):
                        result.extend(d for d in parsed if isinstance(d, dict))
                except json.JSONDecodeError:
                    continue
        return result

    nfrs_list = _ensure_dict_list(data.get('nfrs', []))
    ts_list = _ensure_dict_list(data.get('test_scenarios', []))

    nfrs_by_us = {}
    for nfr in nfrs_list:
        us_id = nfr.get('source_user_story_id', 'UNKNOWN')
        nfrs_by_us.setdefault(us_id, []).append(nfr)

    scenarios_by_nfr = {}
    for ts in ts_list:
        nfr_id = ts.get('source_nfr_id', 'UNKNOWN')
        scenarios_by_nfr.setdefault(nfr_id, []).append(ts)

    # Process each User Story
    user_stories = data.get('user_stories', [])
    for us_index, us in enumerate(user_stories):
        # Heading
        heading = doc.add_heading(f"[{us['id']}]: {us.get('story', us.get('title', ''))}", level=2)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_paragraph_shading(heading, '2F2F2F')
        for run in heading.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.font.bold = True

        # Acceptance Criteria
        ac = us.get('acceptance_criteria', [])
        if ac:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            run = p.add_run('Acceptance Criteria:')
            run.font.bold = True
            run.font.italic = True
            run.font.name = 'Arial'
            run.font.size = Pt(11)

            for criteria in ac:
                _create_bullet_paragraph(doc, criteria, level=0)

        # Metadata
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(10)

        run = p.add_run('Priority: ')
        run.font.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(11)

        run = p.add_run(f"{us.get('priority', 'N/A')} | ")
        run.font.name = 'Arial'
        run.font.size = Pt(11)

        run = p.add_run('Story Points: ')
        run.font.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(11)

        run = p.add_run(str(us.get('story_points', 'N/A')))
        run.font.name = 'Arial'
        run.font.size = Pt(11)

        # NFRs
        for nfr in nfrs_by_us.get(us['id'], []):
            _create_bullet_paragraph(
                doc, nfr['description'], level=0,
                bold_prefix=f"[{nfr['id']}] {nfr['title']}: "
            )

            if nfr.get('acceptance_criteria'):
                _create_bullet_paragraph(
                    doc, nfr['acceptance_criteria'], level=1,
                    bold_prefix="Acceptance Criteria: "
                )

            _create_bullet_paragraph(
                doc,
                f"{nfr.get('category', 'General')} | Priority: {nfr.get('priority', 'N/A')}",
                level=1, bold_prefix="Category: "
            )

            # Test Scenarios
            for ts_i, ts in enumerate(scenarios_by_nfr.get(nfr['id'], []), 1):
                p = _create_bullet_paragraph(
                    doc, ts['title'], level=1,
                    bold_prefix=f"Scenario {ts_i} [{ts['id']}]: "
                )
                for run in p.runs[1:]:
                    run.font.bold = True

                _create_bullet_paragraph(doc, ts.get('test_type', ''), level=2, bold_prefix="Test Type: ")
                _create_bullet_paragraph(doc, ts.get('scenario_description', ''), level=2, bold_prefix="Conditions: ")
                _create_bullet_paragraph(doc, ts.get('expected_outcome', ''), level=2, bold_prefix="Expected Outcome: ")
                _create_bullet_paragraph(doc, ts.get('priority', ''), level=2, bold_prefix="Priority: ")

        # Separator
        if us_index < len(user_stories) - 1:
            p = doc.add_paragraph('_' * 79)
            p.paragraph_format.space_before = Pt(15)
            p.paragraph_format.space_after = Pt(15)
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(11)

    doc.save(output_path)
    return output_path


class DocGeneratorTool(BaseTool):
    """CrewAI tool for generating .docx requirements document."""
    name: str = "Generate Requirements Document"
    description: str = (
        "Generates a .docx requirements document from the complete dataset. "
        "Call this after all user stories, NFRs, and test scenarios have been saved."
    )

    def _run(self) -> str:
        try:
            output_dir = get_output_dir()
            dataset_path = output_dir / "complete_dataset.json"

            if not dataset_path.exists():
                return "Error: complete_dataset.json not found. Save all results first."

            with open(dataset_path, 'r', encoding='utf-8') as f:
                data = json.load(f)

            llm_cfg = get_config().get("llm", {})
            model_name = llm_cfg.get("model", "CrewAI Agents")

            output_path = str(output_dir / "documento_requisitos.docx")
            generate_requirements_doc(data, output_path, model_name)

            return f"Successfully generated requirements document: {output_path}"
        except Exception as e:
            return f"Error generating document: {e}"
